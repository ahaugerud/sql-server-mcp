from datetime import datetime

import pyarrow.parquet as pq

from sql_server_mcp.export import write_csv, write_parquet


def test_write_parquet_roundtrip(tmp_path):
    rows = [
        {"id": 1, "name": "a", "ts": datetime(2026, 1, 1), "amt": None},
        {"id": 2, "name": "b", "ts": datetime(2026, 1, 2), "amt": 1.5},
    ]
    path = tmp_path / "out.parquet"

    write_parquet(rows, str(path))

    assert pq.read_table(str(path)).to_pylist() == rows


def test_write_csv_roundtrip(tmp_path):
    rows = [{"id": "1", "name": "a"}, {"id": "2", "name": "b"}]
    path = tmp_path / "out.csv"

    write_csv(rows, str(path))

    content = path.read_text()
    assert content.splitlines() == ["id,name", "1,a", "2,b"]


def test_write_csv_is_utf8_regardless_of_locale(tmp_path, monkeypatch):
    monkeypatch.setattr("locale.getpreferredencoding", lambda *a, **k: "cp1252")
    rows = [{"name": "Å☃"}]
    path = tmp_path / "out.csv"

    write_csv(rows, str(path))

    assert path.read_bytes().decode("utf-8").splitlines()[1] == "Å☃"


def test_write_excel_roundtrip(tmp_path):
    from datetime import timezone

    from openpyxl import load_workbook

    from sql_server_mcp.export import write_excel

    rows = [
        {"id": 1, "name": "Å\x00b", "ts": datetime(2026, 1, 1, tzinfo=timezone.utc)},
        {"id": 2, "name": None, "ts": None},
    ]
    path = tmp_path / "out.xlsx"

    write_excel(rows, str(path))

    ws = load_workbook(path).active
    assert [[c.value for c in r] for r in ws.iter_rows()] == [
        ["id", "name", "ts"],
        [1, "Åb", datetime(2026, 1, 1)],
        [2, None, None],
    ]


def test_write_excel_rejects_too_many_rows(tmp_path, monkeypatch):
    import pytest

    from sql_server_mcp import export

    monkeypatch.setattr(export, "EXCEL_MAX_ROWS", 1)
    with pytest.raises(ValueError, match="exceeds Excel"):
        export.write_excel([{"a": 1}, {"a": 2}], str(tmp_path / "x.xlsx"))
